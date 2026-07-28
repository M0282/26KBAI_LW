from io import BytesIO

import pytest
from docx import Document

from src.parser.document_loader import (
    InvalidDocumentError,
    TextDocument,
    UnsupportedDocumentError,
    load_document,
    load_docx,
    load_txt,
    to_parsed_document,
)


def docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("적합성 진단표")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "투자성향"
    table.cell(0, 1).text = "안정형"
    table.cell(1, 0).text = "고객 확인"
    table.cell(1, 1).text = "서명 완료"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_loader_preserves_paragraph_and_table_order():
    loaded = load_docx(docx_bytes(), file_name="fit.docx")

    assert loaded.source_format == "docx"
    assert loaded.document_id == "fit.docx"
    assert loaded.blocks == [
        "적합성 진단표",
        "투자성향\t안정형",
        "고객 확인\t서명 완료",
    ]
    assert "투자성향\t안정형" in loaded.text


def test_docx_locator_returns_grounded_block_without_fake_coordinates():
    loaded = load_docx(docx_bytes(), file_name="fit.docx")

    hits = loaded.locate("투자 성향 안정형")

    assert hits == [{
        "page": 1,
        "block": 2,
        "rects": [],
        "excerpt": "투자성향 안정형",
        "location_type": "block",
    }]


@pytest.mark.parametrize("encoding", ["utf-8-sig", "utf-16", "cp949"])
def test_txt_loader_supports_common_korean_encodings(encoding):
    raw = "상품설명서\n위험등급: 1등급".encode(encoding)

    loaded = load_txt(raw, file_name="product.txt")

    assert loaded.blocks == ["상품설명서", "위험등급: 1등급"]
    assert loaded.locate("위험 등급: 1등급")[0]["block"] == 2


def test_duplicate_streamlit_name_keeps_real_extension():
    loaded = load_document(docx_bytes(), file_name="fit.docx (2)")
    assert loaded.document_id == "fit.docx (2)"
    assert loaded.source_format == "docx"


def test_to_parsed_document_keeps_shared_schema_contract():
    loaded = TextDocument(
        document_id="fit.txt",
        blocks=["투자성향: 안정형"],
        source_format="txt",
    )

    parsed = to_parsed_document(loaded)

    assert parsed.document_id == "fit.txt"
    assert parsed.doc_type == "unknown"
    assert parsed.fields == []
    assert parsed.raw_text == "투자성향: 안정형"


def test_invalid_and_unsupported_inputs_are_explicit():
    with pytest.raises(InvalidDocumentError):
        load_txt(b"\n\n", file_name="empty.txt")
    with pytest.raises(UnsupportedDocumentError):
        load_document(b"data", file_name="sample.hwp")


def test_docx_flows_into_existing_field_extractor():
    from src.parser.financial_extractor import extract_rule_based

    loaded = load_docx(docx_bytes(), file_name="fit.docx")
    result = extract_rule_based(to_parsed_document(loaded), locator=loaded.locate)
    values = {field.name: field.value for field in result.fields}

    assert result.doc_type == "suitability_form"
    assert values["customer_profile"] == "안정형"


def test_txt_risk_grade_flows_into_existing_field_extractor():
    from src.parser.financial_extractor import extract_rule_based

    loaded = load_txt(
        "상품설명서\n위험등급: 1등급".encode("utf-8"),
        file_name="product.txt",
    )
    result = extract_rule_based(to_parsed_document(loaded), locator=loaded.locate)
    values = {field.name: field.value for field in result.fields}

    assert result.doc_type == "product_description"
    assert values["product_risk_level"] == "1등급"
